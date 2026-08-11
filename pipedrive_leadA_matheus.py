#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
==============================================================================
 LEAD A - closer MATHEUS MEDEIROS  ->  data + perfil do lead + produto apresentado
==============================================================================

O QUE FAZ
---------
1) Busca a lista de campos (dealFields) uma vez e monta o dicionario
   opcao_id -> texto, para decodificar os campos enum/set (senao eles vem
   como numero).
2) Para cada deal da lista TARGETS, pega o negocio em /deals/{id} e extrai:
       - DATA .............. Dia da Reuniao / Dia Oportunidade / Ganho em / Criado em
       - PERFIL DO LEAD .... grade Lead (A/B/C...), Score, Faturamentos,
                             Onde vende, Estado, Cidade, Cargo, Categoria,
                             Tipo de Negocio, Marketplaces, Tempo de ML,
                             Principal Dor, Objetivo, Pessoa, Organizacao
       - PRODUTO APRESENTADO  campo "Produto Apresentado" (id 280)
3) Confere se a etiqueta "Matheus Medeiros" (id 660) esta no negocio.
4) Salva  leadA_matheus.csv  e  leadA_matheus.json  na pasta atual.

COMO USAR
---------
    export PIPEDRIVE_TOKEN="seu_token_aqui"      # (Windows: set PIPEDRIVE_TOKEN=...)
    python pipedrive_leadA_matheus.py

    # ou passando ids diferentes:
    python pipedrive_leadA_matheus.py 114715 114805 ...
==============================================================================
"""

import os
import sys
import csv
import json
import time

import requests

try:
    from docx import Document
    from docx.shared import Pt
    TEM_DOCX = True
except ImportError:
    TEM_DOCX = False

# =============================================================================
# CONFIG
# =============================================================================
TOKEN = os.environ.get("PIPEDRIVE_TOKEN", "").strip()
COMPANY_DOMAIN = os.environ.get("PIPEDRIVE_DOMAIN", "").strip()  # opcional
PAUSA = 0.15  # segundos entre chamadas (educado com a API)

# etiqueta que marca o closer
CLOSER_LABEL_ID = 660          # "Matheus Medeiros"
CLOSER_LABEL_NOME = "Matheus Medeiros"

# os deals a puxar (pode sobrescrever pela linha de comando)
TARGETS_PADRAO = [114715, 114805, 114841, 114880, 115051, 115054, 115453, 115645, 115026]

# campos que compoem o "perfil do lead" (nome_amigavel -> key da API)
PERFIL_FIELDS = [
    ("lead_grade",              "1cb9ad4729d5f241f8a329045613582070ee2c95"),  # Lead A/B/C...
    ("score",                   "score"),
    ("faturamento_marketplaces","e5310680e97128cdd6fd135118b1c277f0538e21"),
    ("faturamento_fisico",      "e4b95213fed643447308e95399bbff05a7d31d87"),
    ("onde_vende",              "d166e9a3c962c7dcbf849b78100083ca42323505"),
    ("estado",                  "cbef5dcdc3dcb5fd20bb243a6f4291093cea5ed7"),
    ("cidade",                  "c3b792625224f2a850ee2d68a2ed8fb2243b2dd8"),
    ("cargo",                   "2c0d5e90bfdb3ced22e92a136f329b576d2ab2d9"),
    ("categoria_principal",     "4feca3e670cd5fa48432231cde645bca827551de"),
    ("tipo_negocio",            "d020ce7781ab8bebdd90fcf1c7ed874f0457ae64"),
    ("marketplaces_vende",      "85835ceac34a89350688bb976331fbd42f6bbbd7"),
    ("tempo_mercado_livre",     "18b907895a6e7674970f1901d4b81d00e295f8df"),
    ("principal_dor",           "88dbdbdebd39d50e0c4bad8e6cc7c05e3e1e0731"),
    ("objetivo",                "c99ef8ae335443d128839b8b63e3f35c0580125b"),
]

# datas relevantes (nome_amigavel -> key)
DATA_FIELDS = [
    ("dia_reuniao",     "c8a99e668e9913e2362556c43a03bd821db81006"),
    ("dia_oportunidade","395bf927e670b580aa5d012e9f242defca9f3050"),
    ("ganho_em",        "won_time"),
    ("criado_em",       "add_time"),
]

# produto apresentado (id 280) + produto vendido (id 65) para referencia
PRODUTO_APRESENTADO_KEY = "810aeffbcd6e447f89e9b82d1a977738a10f2f86"
PRODUTO_VENDIDO_KEY = "6621eda2196194ded0bd671ae263ccbcb213666c"


# =============================================================================
# API
# =============================================================================
def base_url():
    if COMPANY_DOMAIN:
        return f"https://{COMPANY_DOMAIN}.pipedrive.com/api/v1"
    return "https://api.pipedrive.com/v1"


def api_get(endpoint, params=None, tentativas=3):
    params = dict(params or {})
    params["api_token"] = TOKEN
    url = f"{base_url()}/{endpoint}"
    for t in range(tentativas):
        try:
            r = requests.get(url, params=params,
                             headers={"Accept": "application/json"}, timeout=60)
            if r.status_code == 200:
                return r.json()
            if r.status_code == 429:  # rate limit
                time.sleep(2 * (t + 1))
                continue
            print(f"    !! HTTP {r.status_code} em {endpoint}: {r.text[:150]}")
            return {}
        except requests.RequestException as e:
            if t == tentativas - 1:
                print(f"    !! erro em {endpoint}: {e}")
                return {}
            time.sleep(1.5 * (t + 1))
    return {}


# =============================================================================
# DECODIFICACAO DE CAMPOS
# =============================================================================
def carregar_mapa_opcoes():
    """Monta { field_key: { str(option_id): label } } a partir de dealFields."""
    data = api_get("dealFields", {"limit": 500})
    mapa = {}
    for f in (data.get("data") or []):
        key = f.get("key")
        opts = f.get("options")
        if key and opts:
            mapa[key] = {str(o.get("id")): o.get("label") for o in opts}
    return mapa


def carregar_etapas():
    """id da etapa -> nome (para saber 'onde esta' o negocio no funil)."""
    return {s.get("id"): s.get("name")
            for s in (api_get("stages", {"limit": 500}).get("data") or [])}


def decodificar(valor, key, mapa):
    """Traduz o valor de um campo enum/set para texto; senao devolve como esta."""
    if valor is None or valor == "":
        return ""
    # pessoa/org vem como dict {name: ...}
    if isinstance(valor, dict):
        return valor.get("name") or valor.get("value") or ""
    opcoes = mapa.get(key)
    if not opcoes:
        return valor
    # set (multiplo) pode vir como "1,2,3" ou lista
    if isinstance(valor, str) and "," in valor:
        partes = [p.strip() for p in valor.split(",")]
        return ", ".join(opcoes.get(p, p) for p in partes)
    if isinstance(valor, list):
        return ", ".join(opcoes.get(str(p), str(p)) for p in valor)
    return opcoes.get(str(valor), valor)


def salvar_docx(linhas, cols, caminho="leadA_matheus.docx"):
    """Gera um Word com uma secao por deal (tabela chave/valor)."""
    if not TEM_DOCX:
        print("!! python-docx nao instalado: pule com 'pip install python-docx' "
              "para gerar o Word.")
        return
    doc = Document()
    doc.add_heading("Lead A - Closer Matheus Medeiros", level=0)
    doc.add_paragraph(f"{len(linhas)} negocios | gerado do Pipedrive")
    for l in linhas:
        doc.add_heading(f"deal/{l.get('deal_id','')} - {l.get('titulo','')}", level=1)
        if "erro" in l:
            doc.add_paragraph("NAO ENCONTRADO")
            continue
        t = doc.add_table(rows=0, cols=2)
        t.style = "Light Grid Accent 1"
        for c in cols:
            if c == "deal_id":
                continue
            valor = l.get(c, "")
            if valor in ("", None):
                continue
            row = t.add_row().cells
            row[0].text = c.replace("_", " ").title()
            row[1].text = str(valor)
        doc.add_paragraph("")
    doc.save(caminho)
    print(f"Word: {caminho}")


# =============================================================================
# CLASSIFICACAO OFICIAL (deck "Definicao Classificacao de Leads V5_1")
# Regra: cortes (fat nulo->F; estado NVA->E; cargo<=Supervisor->E). Depois calcula
# a nota pela matriz de Marketplaces (fat online) E pela de Fisico (fat fisico) e
# fica com a MAIOR das duas (leads hibridos usam a melhor classificacao).
# =============================================================================
# matriz Marketplaces: 10 faixas (0-5k, 5-15k, 15-40k, 40-60k, 60-100k, 100-300k,
#                                 300-500k, 500k-1M, 1M-5M, >5M)
_MKT = {"S": ["E", "E", "C", "B", "B", "A", "A", "A", "A", "A"],
        "G": ["E", "E", "E", "E", "D", "D", "D", "D", "D", "D"]}
# matriz Fisico: 8 faixas (0-15k, 15-60k, 60-100k, 100-300k, 300-500k, 500k-1M, 1M-5M, >5M)
_FIS = {"S": ["E", "E", "E", "B", "A", "A", "A", "A"],
        "G": ["E", "E", "E", "D", "D", "D", "D", "D"]}
_ESTADOS_VAL = {"são paulo", "minas gerais", "paraná", "rio grande do sul",
                "santa catarina", "distrito federal", "rio de janeiro"}
_CARGO_S = {"sócio ou fundador", "diretor"}
_CARGO_G = {"gerente", "coordenador"}
_ORD = {"A": 0, "B": 1, "C": 2, "D": 3, "E": 4, "F": 5}


def _tier_mkt(v):
    if not v:
        return None
    v = v.lower()
    if "ainda não" in v or "não temos" in v or "não vendemos" in v:
        return None
    if "acima de r$ 5" in v: return 9
    if "1 milhão a r$ 5" in v: return 8
    if "500.000 a r$ 1" in v: return 7
    if "300.000 a r$ 500" in v: return 6
    if "100.000 a r$ 300" in v: return 5
    if "60.000 a r$ 100" in v: return 4
    if "40.000 a r$ 60" in v: return 3
    if "15.000 a r$ 40" in v or "15.000 a r$ 60" in v: return 2   # faixa larga -> menor
    if "5.000 a r$ 15" in v: return 1
    if "0 a r$ 5.000" in v or "0 a r$ 15" in v: return 0
    return None


def _tier_fis(v):
    if not v:
        return None
    v = v.lower()
    if "ainda não" in v or "não temos" in v or "não vendemos" in v:
        return None
    if "acima de r$ 5" in v: return 7
    if "1 milhão a r$ 5" in v: return 6
    if "500.000 a r$ 1" in v: return 5
    if "300.000 a r$ 500" in v: return 4
    if "100.000 a r$ 300" in v: return 3
    if "60.000 a r$ 100" in v: return 2
    if "15.000 a r$ 60" in v or "15.000 a r$ 40" in v or "40.000 a r$ 60" in v: return 1
    if "0 a r$ 15" in v or "0 a r$ 5.000" in v or "5.000 a r$ 15" in v or "0 a r$ 60" in v: return 0
    return None


def classificar_deck(linha):
    """Retorna (grade, detalhe) pela regra do deck V5_1 (maior nota entre os canais)."""
    c = (linha.get("cargo") or "").lower()
    grupo = "S" if c in _CARGO_S else ("G" if c in _CARGO_G else None)
    if grupo is None:
        return "E", "cargo <= Supervisor"
    if (linha.get("estado") or "").lower() not in _ESTADOS_VAL:
        return "E", "estado nao valido"
    tm = _tier_mkt(linha.get("faturamento_marketplaces"))
    tf = _tier_fis(linha.get("faturamento_fisico"))
    if tm is None and tf is None:
        return "F", "nao vende em nenhum canal"
    g_mkt = _MKT[grupo][tm] if tm is not None else None
    g_fis = _FIS[grupo][tf] if tf is not None else None
    grade = min([g for g in (g_mkt, g_fis) if g], key=lambda g: _ORD[g])  # MAIOR nota
    det = (f"online->{g_mkt}" if g_mkt else "online=nao vende")
    det += (f" | fisico->{g_fis}" if g_fis else " | fisico=nao vende")
    det += f" => maior={grade}"
    return grade, det


def tem_etiqueta_closer(deal):
    """True se a etiqueta Matheus Medeiros (660) esta no negocio."""
    ids = deal.get("label_ids")
    if isinstance(ids, list):
        return CLOSER_LABEL_ID in [int(x) for x in ids if str(x).isdigit()]
    lab = deal.get("label")
    if lab is None:
        return False
    partes = str(lab).split(",")
    return str(CLOSER_LABEL_ID) in [p.strip() for p in partes]


# =============================================================================
# MAIN
# =============================================================================
def main():
    if not TOKEN:
        print("!! Defina o token:  export PIPEDRIVE_TOKEN=...  (ou set no Windows)")
        return

    targets = [int(a) for a in sys.argv[1:] if a.isdigit()] or TARGETS_PADRAO
    print(f"Puxando {len(targets)} deals...")

    mapa = carregar_mapa_opcoes()
    if not mapa:
        print("!! Nao consegui carregar dealFields (token/conexao?). Abortando.")
        return
    etapas = carregar_etapas()

    linhas = []
    for did in targets:
        d = api_get(f"deals/{did}").get("data")
        if not d:
            print(f"  deal/{did}: NAO ENCONTRADO")
            linhas.append({"deal_id": did, "erro": "nao encontrado"})
            time.sleep(PAUSA)
            continue

        linha = {
            "deal_id": did,
            "titulo": d.get("title", ""),
            "status": d.get("status", ""),
            "valor": d.get("value", ""),
            "etiqueta_matheus": "SIM" if tem_etiqueta_closer(d) else "NAO",
            "etapa": etapas.get(d.get("stage_id"), d.get("stage_id") or ""),
            # "Motivo da Perda" (dropdown) — costuma vir vazio
            "motivo_perdido": decodificar(d.get("lost_reason"), "lost_reason", mapa),
            # "Descricao DETALHADA da Perda" (texto livre, id 282) — o motivo real
            "descricao_perda": (d.get("c1056dfb2bc711a3a2dc66b4dc9c83a436a36a84") or "").replace("\n", " ").strip(),
        }
        # datas
        for nome, key in DATA_FIELDS:
            linha[nome] = d.get(key) or ""
        # produto apresentado / vendido
        linha["produto_apresentado"] = decodificar(d.get(PRODUTO_APRESENTADO_KEY),
                                                    PRODUTO_APRESENTADO_KEY, mapa)
        linha["produto_vendido"] = decodificar(d.get(PRODUTO_VENDIDO_KEY),
                                                PRODUTO_VENDIDO_KEY, mapa)
        # perfil do lead
        for nome, key in PERFIL_FIELDS:
            linha[nome] = decodificar(d.get(key), key, mapa)
        # pessoa / organizacao
        linha["pessoa"] = decodificar(d.get("person_id"), "person_id", mapa)
        linha["organizacao"] = decodificar(d.get("org_id"), "org_id", mapa)
        # auditoria: grade pelo deck V5 e se bate com o campo Lead do pipe
        deck, det = classificar_deck(linha)
        linha["deck_grade"] = deck
        linha["deck_detalhe"] = det
        linha["bate_A"] = "SIM" if (linha.get("lead_grade") == "A" and deck == "A") else \
                          ("NAO" if linha.get("lead_grade") == "A" else "")

        marca = "OK" if linha["etiqueta_matheus"] == "SIM" else "(sem etiqueta Matheus)"
        print(f"  deal/{did}: {linha['titulo'][:40]:<40} "
              f"[{linha['status']}] {marca}")
        linhas.append(linha)
        time.sleep(PAUSA)

    # ordem das colunas
    cols = ["deal_id", "titulo", "status", "etapa", "motivo_perdido", "descricao_perda",
            "valor", "etiqueta_matheus"]
    cols += [n for n, _ in DATA_FIELDS]
    cols += ["produto_apresentado", "produto_vendido"]
    cols += [n for n, _ in PERFIL_FIELDS]
    cols += ["deck_grade", "bate_A", "deck_detalhe"]
    cols += ["pessoa", "organizacao"]

    with open("leadA_matheus.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for l in linhas:
            w.writerow(l)
    with open("leadA_matheus.json", "w", encoding="utf-8") as f:
        json.dump(linhas, f, ensure_ascii=False, indent=2)
    salvar_docx(linhas, cols)

    achou = sum(1 for l in linhas if "erro" not in l)
    com_etq = sum(1 for l in linhas if l.get("etiqueta_matheus") == "SIM")
    print("-" * 64)
    print(f"Deals encontrados: {achou}/{len(targets)}")
    print(f"Com etiqueta Matheus Medeiros: {com_etq}")
    print("Arquivos: leadA_matheus.csv  |  leadA_matheus.json")


if __name__ == "__main__":
    main()
